#!/usr/bin/env python3
"""
exportar_correlacoes_vin.py — gera `correlacoes_vin.xlsx`: só os cruzamentos entre as tabelas feitos pelo VIN
(DPHU × GMF, DPHU × AVES, AVES × GMF) e a lista de VINs ligados. Os números são os mesmos do analise_dados.py.
"""
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from scipy import stats

import analise_dados as A
import treinar as T

AQUI = Path(__file__).resolve().parent
SAIDA = AQUI / "correlacoes_vin.xlsx"
FAMS = list("GRBCIHOVESA")


def tabela_sir(d, col, claim, minn):
    r = A.sir_tabela(d, col, claim, minn).reset_index()
    r.columns = ["Grupo", "VINs no DPHU", "Com reclamação (observado)", "Esperado", "Observado ÷ esperado", "p", "q (corrigido)"]
    r["Leitura"] = np.where(r["q (corrigido)"] < 0.05, np.where(r["Observado ÷ esperado"] > 1, "ACIMA do esperado", "ABAIXO do esperado"), "dentro do esperado")
    return r


def pares_familia(d, g):
    d = d.assign(dfam=d["Final Element Label"].map({l: T.depara(l) or "?" for l in d["Final Element Label"].unique()}), strata=d.mes + "|" + d["mod"])
    gg = g[g.VIN.isin(d.VIN) & g.fam.isin(FAMS) & (g.cod != "G057")][["VIN", "fam"]].drop_duplicates()
    Y = pd.DataFrame(0, index=d.VIN.values, columns=FAMS)
    for v, f in gg.itertuples(index=False):
        Y.at[v, f] = 1
    R = Y.groupby(d.strata.values).transform("mean").values; Yv = Y.values; rows = []
    for e in d.dfam.unique():
        m = (d.dfam == e).values
        if m.sum() < 200:
            continue
        O, E = Yv[m].sum(0), R[m].sum(0)
        for k, f in enumerate(FAMS):
            if O[k] >= 5:
                p = min(1, 2 * (stats.poisson.sf(O[k] - 1, E[k]) if O[k] >= E[k] else stats.poisson.cdf(O[k], E[k])))
                rows.append((e, f, int(m.sum()), int(O[k]), E[k], O[k] / E[k], p))
    r = pd.DataFrame(rows, columns=["Família do defeito de fábrica (estimada pelo nome do defeito)", "Família da reclamação (GMF)", "VINs no DPHU",
                                    "Com reclamação (observado)", "Esperado", "Observado ÷ esperado", "p"])
    r["q (corrigido)"] = A.bh(r.p)
    r["Mesmo sistema?"] = np.where(r.iloc[:, 0] == r.iloc[:, 1], "sim", "")
    return r.sort_values("p")


def correlacao_lotes(d, g):
    d = d.assign(dfam=d["Final Element Label"].map({l: T.depara(l) or "?" for l in d["Final Element Label"].unique()}))
    gl = g[g.cod.notna() & (g.cod != "G057")].copy()
    gl["lote"] = gl["MODEL CODE"] + "|" + gl.mfg.dt.to_period("W").astype(str)
    d["lote"] = d["mod"] + "|" + d["Vehicle MADU Date"].dt.to_period("W").astype(str)
    nd, ng = d.groupby("lote").size(), gl.groupby("lote").size()
    lotes = [l for l in nd.index if l in ng.index and nd[l] >= 100 and ng[l] >= 30]
    out = []
    for f in "GRCIHOVE":
        x = d[d.dfam == f].groupby("lote").size().reindex(lotes).fillna(0) / nd[lotes]
        y = gl[gl.fam == f].groupby("lote").size().reindex(lotes).fillna(0) / ng[lotes]
        rho, p = stats.spearmanr(x, y)
        out.append((f, len(lotes), rho, p, min(1, p * 8)))
    r = pd.DataFrame(out, columns=["Família", "Lotes (modelo × semana)", "Correlação (Spearman)", "p", "p corrigido (×8 famílias)"])
    r["Leitura"] = np.where(r["p corrigido (×8 famílias)"] < 0.05, "relação positiva forte o bastante", "sem relação clara")
    return r


def vins_ligados(d, a, g):
    a1 = a.dropna(subset=["VIN"])
    ids = (set(d.VIN) & set(g.VIN)) | (set(d.VIN) & set(a1.VIN)) | (set(a1.VIN) & set(g.VIN))
    dd = d[d.VIN.isin(ids)].set_index("VIN")
    aa = a1[a1.VIN.isin(ids)].groupby("VIN").agg(n_aves=("VIN", "size"), aves_demerito=("Demérito (pt)", lambda s: ", ".join(sorted(set(s.str.replace("(", "").str.replace(")", ""))))),
                                                  aves_defeitos=("Descrição defeito", lambda s: " | ".join(s.head(3))))
    gc = g[g.VIN.isin(ids)].groupby("VIN").agg(n_gmf=("VIN", "size"), gmf_modelo=("MODEL CODE", "first"), gmf_mes_fab=("mfg", lambda s: s.iloc[0].strftime("%Y-%m")),
                                                gmf_codigos=("cod", lambda s: ", ".join(dict.fromkeys(s.dropna()))),
                                                gmf_primeira_abertura=("cr", "min"), gmf_km_primeira=("MILEAGE (KM)", "first"))
    t = pd.DataFrame(index=sorted(ids))
    t["No DPHU"] = ["sim" if v in dd.index else "" for v in t.index]
    t["No AVES"] = ["sim" if v in aa.index else "" for v in t.index]
    t["No GMF"] = ["sim" if v in gc.index else "" for v in t.index]
    t["Modelo"] = [dd["mod"].get(v) if v in dd.index else (gc["gmf_modelo"].get(v) if v in gc.index else None) for v in t.index]
    t["Mês de fabricação"] = [dd["mes"].get(v) if v in dd.index else (gc["gmf_mes_fab"].get(v) if v in gc.index else None) for v in t.index]
    t["DPHU · elemento"] = dd["elem"].reindex(t.index).values
    t["DPHU · incidente"] = dd["inc"].reindex(t.index).values
    t["DPHU · demérito"] = dd["Demerit"].fillna("sem demérito").reindex(t.index).values
    t["AVES · nº de defeitos"] = aa["n_aves"].reindex(t.index).values
    t["AVES · demérito"] = aa["aves_demerito"].reindex(t.index).values
    t["AVES · defeitos (até 3)"] = aa["aves_defeitos"].reindex(t.index).values
    t["GMF · nº de casos"] = gc["n_gmf"].reindex(t.index).values
    t["GMF · códigos de sintoma"] = gc["gmf_codigos"].reindex(t.index).values
    t["GMF · 1ª abertura"] = gc["gmf_primeira_abertura"].reindex(t.index).dt.date.values
    t["GMF · km na 1ª abertura"] = gc["gmf_km_primeira"].reindex(t.index).values
    t.index.name = "VIN"
    return t.reset_index()


def main():
    d, a, g, desc = A.carregar()
    d["Demerit"] = d.Demerit  # mantém NaN para o resumo
    com_dg, com_da, com_ag = set(d.VIN) & set(g.VIN), set(d.VIN) & set(a.VIN.dropna()), set(a.VIN.dropna()) & set(g.VIN)
    tri = com_dg & com_da & com_ag
    j = d.merge(g.drop_duplicates("VIN")[["VIN", "MODEL CODE", "mfg"]], on="VIN")
    jd = a.dropna(subset=["VIN"]).merge(d[["VIN", "Type Code", "Vehicle MADU Date", "Declaration Date Time"]], on="VIN")
    ga = g.merge(a.dropna(subset=["VIN"]).drop_duplicates("VIN")[["VIN", "dt", "Tipo de veículo"]], on="VIN")
    gd = g[g.VIN.isin(com_dg)].merge(d[["VIN", "Vehicle MADU Date"]], on="VIN")
    resumo = pd.DataFrame([
        ["DPHU × GMF", len(com_dg), len(com_dg) / d.VIN.nunique(), len(com_dg) / g.VIN.nunique(),
         (j["mod"] == j["MODEL CODE"]).mean(), (j["Vehicle MADU Date"].dt.to_period("M") == j.mfg.dt.to_period("M")).mean(),
         "Mediana de %d dias da fabricação até a abertura do caso" % (gd.cr - gd["Vehicle MADU Date"]).dt.days.median()],
        ["DPHU × AVES", len(com_da), len(com_da) / d.VIN.nunique(), len(com_da) / a.VIN.nunique(),
         (jd["Tipo de veículo"] == jd["Type Code"]).mean(), (jd.dt.dt.to_period("M") == jd["Vehicle MADU Date"].dt.to_period("M")).mean(),
         "Auditoria 1 dia depois do defeito declarado (mediana)"],
        ["AVES × GMF", len(com_ag), len(com_ag) / a.VIN.nunique(), len(com_ag) / g.VIN.nunique(),
         (ga["Tipo de veículo"].dropna() == ga.loc[ga["Tipo de veículo"].notna(), "MODEL CODE"]).mean(), np.nan,
         "Mediana de %d dias da auditoria até a abertura do caso" % (ga.cr - ga.dt).dt.days.median()],
        ["DPHU × AVES × GMF", len(tri), np.nan, np.nan, np.nan, np.nan, "Pequeno demais para concluir"],
    ], columns=["Cruzamento", "VINs em comum", "% dos VINs da 1ª tabela", "% dos VINs da 2ª tabela", "Mesmo modelo", "Mesmo mês de fabricação", "Observação"])
    resumo.loc[2, "% dos VINs da 1ª tabela"] = np.nan
    # ajuste das porcentagens: 1ª tabela = a primeira citada no nome
    resumo.loc[1, ["% dos VINs da 1ª tabela", "% dos VINs da 2ª tabela"]] = [len(com_da) / d.VIN.nunique(), len(com_da) / a.VIN.nunique()]
    resumo.loc[2, ["% dos VINs da 1ª tabela", "% dos VINs da 2ª tabela"]] = [len(com_ag) / a.VIN.nunique(), len(com_ag) / g.VIN.nunique()]

    claim = set(g[g.cod != "G057"].VIN)
    d2 = d.assign(Demerit=d.Demerit.fillna("sem demérito"))
    abas = {
        "Resumo": resumo,
        "DPHU×GMF · demérito": tabela_sir(d2, "Demerit", claim, 5),
        "DPHU×GMF · incidente": tabela_sir(d, "inc", claim, 100),
        "DPHU×GMF · elemento": tabela_sir(d, "elem", claim, 100),
        "DPHU×GMF · família": pares_familia(d, g),
        "DPHU×GMF · lotes": correlacao_lotes(d, g),
        "VINs ligados": vins_ligados(d, a, g),
    }
    notas = [
        ["Leitura das abas", ""],
        ["Observado ÷ esperado", "Reclamações (sem G057) entre os VINs do grupo ÷ reclamações esperadas, descontando mês de fabricação e modelo. 1,00 = igual ao esperado; acima = mais reclamação que o esperado."],
        ["p e q", "p = chance de acaso; q = p corrigido para muitos testes (Benjamini-Hochberg). Só q < 0,05 é tratado como diferença real."],
        ["Por que descontar mês e modelo", "Carros fabricados antes tiveram mais tempo para reclamar. Sem o desconto, o grupo mais antigo pareceria pior."],
        ["Família do defeito de fábrica", "É estimada pelas palavras do nome do defeito (não existe tabela de tradução entre as bases)."],
        ["Lotes", "Correlação entre a % de defeitos de fábrica da família e a % de reclamações da família, por modelo × semana de fabricação (lotes com 100+ defeitos DPHU e 30+ casos GMF)."],
        ["Demérito", "V1+ crítico · V1 concessionária imediata · V2 não vai direto · V3 muito leve · 'sem demérito' = testes de função."],
        ["Reclamações", "Todas as colunas de reclamação excluem o código G057 (provável registro de revisão; a confirmar com a Qualidade)."],
        ["Fonte", "Transformation Day Qualidade - Dados 2.xlsx · gerado por exportar_correlacoes_vin.py · dados anonimizados, uso interno."],
    ]
    with pd.ExcelWriter(SAIDA, engine="openpyxl") as w:
        pd.DataFrame(notas).to_excel(w, sheet_name="Como ler", index=False, header=False)
        for nome, df in abas.items():
            df.to_excel(w, sheet_name=nome[:31], index=False)

    # ---------------- formatação
    wb = load_workbook(SAIDA)
    azul, bordo_fundo = "1F4FD8", "F6E7EB"
    fonte = Font(name="Arial", size=10); cab = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    fino = Side(style="thin", color="DFE2EA")
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                c.font = fonte; c.alignment = Alignment(vertical="top", wrap_text=ws.title == "Como ler")
        if ws.title != "Como ler":
            for c in ws[1]:
                c.font = cab; c.fill = PatternFill("solid", fgColor=azul); c.alignment = Alignment(wrap_text=True, vertical="center")
            ws.row_dimensions[1].height = 32; ws.freeze_panes = "A2"
            if ws.max_row > 1:
                ws.auto_filter.ref = ws.dimensions
        for i, col in enumerate(ws.columns, 1):
            largura = max((len(str(c.value)) for c in col[:200] if c.value is not None), default=8)
            ws.column_dimensions[get_column_letter(i)].width = min(60, max(11, largura + 2))
        heads = {c.value: c.column for c in ws[1]}
        for nome_col, fmt in [("Observado ÷ esperado", "0.00"), ("Esperado", "0.0"), ("p", "0.000"), ("q (corrigido)", "0.000"), ("Correlação (Spearman)", "0.00"),
                              ("p corrigido (×8 famílias)", "0.000"), ("% dos VINs da 1ª tabela", "0.0%"), ("% dos VINs da 2ª tabela", "0.0%"),
                              ("Mesmo modelo", "0%"), ("Mesmo mês de fabricação", "0%"), ("VINs em comum", "#,##0"), ("VINs no DPHU", "#,##0")]:
            if nome_col in heads and ws.title != "Como ler":
                for r in range(2, ws.max_row + 1):
                    ws.cell(r, heads[nome_col]).number_format = fmt
        # destaque (bordô claro) nas linhas com diferença real
        q = heads.get("q (corrigido)") or heads.get("p corrigido (×8 famílias)")
        if q:
            for r in range(2, ws.max_row + 1):
                v = ws.cell(r, q).value
                if v is not None and v < 0.05:
                    for c in ws[r]:
                        c.fill = PatternFill("solid", fgColor=bordo_fundo)
    ws = wb["Como ler"]; ws.column_dimensions["A"].width = 34; ws.column_dimensions["B"].width = 110; ws["A1"].font = Font(name="Arial", size=12, bold=True, color=azul)
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 1).font = Font(name="Arial", size=10, bold=True)
    wb.move_sheet("Como ler", offset=-wb.index(wb["Como ler"]))
    wb.save(SAIDA)
    print("gerado:", SAIDA.name, {n: len(df) for n, df in abas.items()})


if __name__ == "__main__":
    main()
